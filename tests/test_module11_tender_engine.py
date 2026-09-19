"""
Module 11 Test Suite — Tender Document Processing / Tender Document Engine.
Tests:
1. Multi-Format Document Parsing:
   - PDF parser using PyMuPDF (pymupdf) with page and text preservation
   - DOCX parser using python-docx with paragraph and table preservation
   - TXT parser with structured section detection
2. Section & Clause Detection:
   - Extraction of section titles, section numbers, and classification of SectionTypes
   - Detection of numbered clauses, mandatory modal verbs, and technical attributes
3. Explicit Standard Reference Extraction & Grounding:
   - Extraction of standard citations (IS 12615, IS 694, IS 269, IS/IEC)
   - Grounding against canonical SQLite standards registry
   - Supersession detection via Module 4 graph (IS 325 -> IS 12615)
   - Verbatim preservation of unlinked/unresolved standard references without data loss
4. Database Persistence:
   - Records created across tender_documents, tender_sections, tender_requirements, and tender_standard_references
5. Section 67 Tender API Endpoints:
   - POST /api/tenders/upload
   - GET /api/tenders/{id}
   - GET /api/tenders/{id}/status
   - GET /api/tenders/{id}/requirements
   - GET /api/tenders/{id}/references
6. Error Handling:
   - Empty files, corrupt payloads, and unsupported file extensions return HTTP 400
7. Health Check:
   - /api/health includes Module 11
"""
import io
import os
import sys
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import sessionmaker

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from backend.app.main import app
from backend.app.core.database import get_engine, get_db
from backend.app.models.standard import Standard
from backend.app.models.tender import (
    TenderDocument,
    TenderSection,
    TenderRequirement,
    TenderStandardReference,
)
from backend.app.services.tender_engine import (
    PDFParser,
    DOCXParser,
    TXTParser,
    get_parser,
    ClauseDetector,
    StandardExtractor,
    TenderEngineService,
    SectionType,
    TenderProcessingStatus,
)


def create_sample_pdf_bytes() -> bytes:
    """Generates an in-memory 2-page PDF document using PyMuPDF."""
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf

    doc = pymupdf.open()
    # Page 1: Scope & General Specifications
    p1 = doc.new_page()
    p1.insert_text((50, 70), "TENDER SPECIFICATION: SUPPLY OF INDUCTION MOTORS\n", fontsize=14)
    p1.insert_text((50, 100), "SECTION 1: SCOPE OF WORK\n", fontsize=12)
    p1.insert_text((50, 130), "Clause 1.1: The scope comprises supply of energy efficient induction motors.\n")
    p1.insert_text((50, 160), "Clause 1.2: All supplied motors shall be 3-phase, 415 V, 50 Hz, 55 kW rating.\n")

    # Page 2: Technical & Compliance Specifications
    p2 = doc.new_page()
    p2.insert_text((50, 70), "SECTION 3: TECHNICAL SPECIFICATIONS AND COMPLIANCE\n", fontsize=12)
    p2.insert_text((50, 100), "Clause 3.1: Motors shall comply with IS 12615:2018 for efficiency classes.\n")
    p2.insert_text((50, 130), "Clause 3.2: General performance shall adhere to IS 325 and IS/IEC 60034-1.\n")
    p2.insert_text((50, 160), "Clause 3.3: Cables used for connection shall conform to IS 694.\n")

    pdf_bytes = doc.write()
    doc.close()
    return pdf_bytes


def create_sample_docx_bytes() -> bytes:
    """Generates an in-memory DOCX document using python-docx."""
    import docx

    doc = docx.Document()
    doc.add_heading("TENDER DOCUMENT: CABLE PROCUREMENT", level=1)
    doc.add_heading("SECTION 1: GENERAL TERMS", level=2)
    doc.add_paragraph("Clause 1.0: Instructions to bidders for procurement of power cables.")
    
    doc.add_heading("SECTION 2: TECHNICAL SPECIFICATIONS", level=2)
    doc.add_paragraph("Clause 2.1: The cables shall be PVC insulated heavy duty cables rated for 1.1 kV.")
    doc.add_paragraph("Clause 2.2: The cables must strictly comply with IS 694 and IS 1554 (Part 1).")
    doc.add_paragraph("Clause 2.3: Cement used for foundation works shall comply with IS 269.")

    # Add a table with parameters
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Standard"
    table.rows[1].cells[0].text = "Electric Motor"
    table.rows[1].cells[1].text = "IS 12615"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_sample_txt_content() -> bytes:
    """Generates sample plain text tender specification."""
    text = (
        "GOVERNMENT OF INDIA - PROCUREMENT TENDER\n\n"
        "SECTION 1: SCOPE OF WORK\n"
        "1.1 Scope includes supply of distribution transformers and switchgear.\n"
        "1.2 All items shall be brand new and manufactured under quality standards.\n\n"
        "SECTION 2: TECHNICAL SPECIFICATIONS\n"
        "Clause 2.1: Motors shall comply with IS 12615:2018 energy efficiency.\n"
        "Clause 2.2: Legacy equipment previously installed under IS 325 shall be upgraded.\n"
        "Clause 2.3: All wiring shall be done using cables conforming to IS 694.\n"
        "Clause 2.4: Civil work cement shall adhere to IS 269.\n"
        "Clause 2.5: Special high-spec component must meet IS 99999 (Unlisted Standard).\n"
    )
    return text.encode("utf-8")


@pytest.fixture(scope="module")
def db_session():
    """Provides a database session over the real canonical database."""
    engine = get_engine("sqlite:///./sih_bis.db")
    Session = sessionmaker(bind=engine)
    session = Session()
    yield session
    session.close()


@pytest.fixture(scope="module")
def client(db_session):
    """Provides a FastAPI TestClient with real database dependency override."""
    app.dependency_overrides[get_db] = lambda: db_session
    client = TestClient(app)
    yield client
    app.dependency_overrides.pop(get_db, None)


# 1. Document Parser Tests
def test_pdf_parser_pymupdf():
    """Test PDF parsing via PyMuPDF preserves pages, text, and metadata."""
    pdf_bytes = create_sample_pdf_bytes()
    parser = PDFParser()
    result = parser.parse(pdf_bytes, filename="tender_motors.pdf")

    assert result.file_type == "PDF"
    assert result.total_pages == 2
    assert len(result.pages) == 2
    assert "INDUCTION MOTORS" in result.pages[0].text
    assert "IS 12615:2018" in result.pages[1].text
    assert result.pages[0].page_number == 1
    assert result.pages[1].page_number == 2


def test_docx_parser():
    """Test DOCX parsing via python-docx extracts paragraphs and tables."""
    docx_bytes = create_sample_docx_bytes()
    parser = DOCXParser()
    result = parser.parse(docx_bytes, filename="cable_tender.docx")

    assert result.file_type == "DOCX"
    assert result.total_pages >= 1
    assert "CABLE PROCUREMENT" in result.raw_text
    assert "IS 694" in result.raw_text
    assert "Electric Motor" in result.raw_text  # Table cell extraction


def test_txt_parser():
    """Test TXT parser extracts text cleanly."""
    txt_bytes = create_sample_txt_content()
    parser = TXTParser()
    result = parser.parse(txt_bytes, filename="tender.txt")

    assert result.file_type == "TXT"
    assert result.total_pages >= 1
    assert "GOVERNMENT OF INDIA" in result.raw_text
    assert "IS 12615:2018" in result.raw_text


def test_parser_factory():
    """Test parser factory correctly identifies file format."""
    assert isinstance(get_parser("doc.pdf"), PDFParser)
    assert isinstance(get_parser("file.docx"), DOCXParser)
    assert isinstance(get_parser("spec.txt"), TXTParser)
    assert isinstance(get_parser("unknown"), TXTParser)


# 2. Clause & Section Detection Tests
def test_clause_detector_sectioning_and_clauses():
    """Test ClauseDetector segments sections, detects clauses, and classifies types."""
    pdf_bytes = create_sample_pdf_bytes()
    parse_result = PDFParser().parse(pdf_bytes)

    detector = ClauseDetector()
    sections = detector.segment_sections_and_clauses(parse_result.pages)

    assert len(sections) >= 2
    # Verify Section 1: Scope of work
    sec1 = sections[0]
    assert "SCOPE OF WORK" in sec1.section_title.upper()
    assert sec1.section_type in [SectionType.SCOPE_OF_WORK, SectionType.TECHNICAL_SPEC]
    assert len(sec1.clauses) >= 1

    # Verify Section 3: Technical specifications
    sec2 = sections[1]
    assert "TECHNICAL SPECIFICATIONS" in sec2.section_title.upper()
    assert sec2.section_type == SectionType.TECHNICAL_SPEC

    # Verify clause detection & technical attributes
    clause_texts = [c.text for c in sec1.clauses]
    has_voltage = any("415" in c.technical_attributes.get("voltage_rating", "") for c in sec1.clauses)
    assert has_voltage or any("415 V" in t for t in clause_texts)


def test_clause_mandatory_detection():
    """Test detection of mandatory requirements via modal verbs."""
    detector = ClauseDetector()
    text_mand = "All equipment shall strictly comply with applicable BIS norms."
    text_non_mand = "The supplier may offer optional accessories if desired."

    is_sec, _, _ = detector.is_section_header(text_mand)
    assert not is_sec
    assert any(kw in text_mand.lower() for kw in detector.MANDATORY_KEYWORDS)
    assert not any(kw in text_non_mand.lower() for kw in detector.MANDATORY_KEYWORDS)


# 3. Standard Extraction & Grounding Tests
def test_standard_extraction_and_grounding(db_session):
    """Test standard extraction, database lookup, and supersession resolution."""
    txt_bytes = create_sample_txt_content()
    parse_result = TXTParser().parse(txt_bytes)
    sections = ClauseDetector().segment_sections_and_clauses(parse_result.pages)

    extractor = StandardExtractor(db_session)
    refs = extractor.extract_from_sections(sections)

    extracted_ids = {r.canonical_identifier for r in refs}
    assert any("12615" in cid for cid in extracted_ids)
    assert any("694" in cid for cid in extracted_ids)
    assert any("269" in cid for cid in extracted_ids)
    assert any("325" in cid for cid in extracted_ids)

    # Check resolution for IS 12615
    ref_12615 = next((r for r in refs if "12615" in r.canonical_identifier), None)
    assert ref_12615 is not None
    assert ref_12615.is_valid is True
    assert ref_12615.detected_standard_id is not None
    assert "motor" in ref_12615.title.lower()

    # Check supersession for IS 325
    ref_325 = next((r for r in refs if "325" in r.canonical_identifier), None)
    assert ref_325 is not None
    assert ref_325.is_superseded is True
    assert ref_325.superseded_by_identifier is not None
    assert "12615" in ref_325.superseded_by_identifier

    # Check unlinked / unlisted standard preservation (IS 99999)
    ref_unlisted = next((r for r in refs if "99999" in r.canonical_identifier), None)
    assert ref_unlisted is not None
    assert ref_unlisted.is_valid is False
    assert ref_unlisted.detected_standard_id is None
    assert "99999" in ref_unlisted.standard_number_raw


# 4. Tender Engine Service & Database Persistence Tests
def test_tender_service_process_document(db_session):
    """Test end-to-end processing and persistence in database."""
    pdf_bytes = create_sample_pdf_bytes()
    service = TenderEngineService(db_session)

    tender_doc = service.process_document(
        file_content=pdf_bytes,
        filename="motor_procurement_2026.pdf",
        title="Supply of Energy Efficient Motors",
        organization="Bharat Heavy Electricals Limited",
    )

    assert tender_doc.id is not None
    assert tender_doc.status == TenderProcessingStatus.COMPLETED.value
    assert tender_doc.file_type == "PDF"
    assert len(tender_doc.sections) >= 2
    assert len(tender_doc.requirements) >= 2

    # Verify query methods
    loaded_tender = service.get_tender(tender_doc.id)
    assert loaded_tender is not None
    assert loaded_tender.title == "Supply of Energy Efficient Motors"

    reqs = service.get_requirements(tender_doc.id)
    assert len(reqs) >= 2

    refs = service.get_references(tender_doc.id)
    assert len(refs) >= 1
    assert any("12615" in r.standard_number_raw for r in refs)


# 5. Section 67 Tender API Endpoints
def test_api_upload_tender_endpoint(client):
    """Test POST /api/tenders/upload endpoint with PDF document."""
    pdf_bytes = create_sample_pdf_bytes()
    response = client.post(
        "/api/tenders/upload",
        files={"file": ("tender_sample.pdf", pdf_bytes, "application/pdf")},
        data={"title": "Induction Motor Procurement", "organization": "Indian Railways"},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["tender_id"] is not None
    assert data["status"] == "COMPLETED"
    assert data["total_pages"] == 2
    assert data["total_sections"] >= 2
    assert data["total_clauses"] >= 2
    assert data["total_standards_detected"] >= 1


def test_api_get_tender_details(client):
    """Test GET /api/tenders/{id}."""
    txt_bytes = create_sample_txt_content()
    up_resp = client.post(
        "/api/tenders/upload",
        files={"file": ("test_tender.txt", txt_bytes, "text/plain")},
        data={"title": "Transformer & Cables Tender"},
    )
    t_id = up_resp.json()["tender_id"]

    resp = client.get(f"/api/tenders/{t_id}")
    assert resp.status_code == 200
    details = resp.json()
    assert details["id"] == t_id
    assert details["status"] == "COMPLETED"
    assert len(details["sections"]) >= 1


def test_api_get_tender_status(client):
    """Test GET /api/tenders/{id}/status."""
    txt_bytes = create_sample_txt_content()
    up_resp = client.post(
        "/api/tenders/upload",
        files={"file": ("status_test.txt", txt_bytes, "text/plain")},
    )
    t_id = up_resp.json()["tender_id"]

    resp = client.get(f"/api/tenders/{t_id}/status")
    assert resp.status_code == 200
    status_data = resp.json()
    assert status_data["tender_id"] == t_id
    assert status_data["status"] == "COMPLETED"
    assert status_data["progress_percentage"] == 100


def test_api_get_tender_requirements(client):
    """Test GET /api/tenders/{id}/requirements."""
    txt_bytes = create_sample_txt_content()
    up_resp = client.post(
        "/api/tenders/upload",
        files={"file": ("reqs_test.txt", txt_bytes, "text/plain")},
    )
    t_id = up_resp.json()["tender_id"]

    resp = client.get(f"/api/tenders/{t_id}/requirements")
    assert resp.status_code == 200
    reqs = resp.json()
    assert reqs["tender_id"] == t_id
    assert reqs["total_requirements"] >= 1
    assert any("IS 12615" in r["requirement_text"] for r in reqs["requirements"])


def test_api_get_tender_references(client):
    """Test GET /api/tenders/{id}/references."""
    txt_bytes = create_sample_txt_content()
    up_resp = client.post(
        "/api/tenders/upload",
        files={"file": ("refs_test.txt", txt_bytes, "text/plain")},
    )
    t_id = up_resp.json()["tender_id"]

    resp = client.get(f"/api/tenders/{t_id}/references")
    assert resp.status_code == 200
    refs = resp.json()
    assert refs["tender_id"] == t_id
    assert refs["total_references"] >= 1
    ref_nums = [r["standard_number_raw"] for r in refs["references"]]
    assert any("12615" in n for n in ref_nums)


def test_api_get_tender_not_found(client):
    """Test GET /api/tenders/999999 returns 404."""
    resp = client.get("/api/tenders/999999")
    assert resp.status_code == 404


def test_api_upload_unsupported_format_returns_400(client):
    """Test uploading unsupported file format (.exe) returns 400 Bad Request."""
    resp = client.post(
        "/api/tenders/upload",
        files={"file": ("malware.exe", b"MZ_binary_executable_data_xyz", "application/octet-stream")},
    )
    assert resp.status_code == 400
    assert "unsupported" in resp.json()["detail"].lower()


def test_api_upload_empty_file_returns_400(client):
    """Test uploading empty file returns 400."""
    resp = client.post(
        "/api/tenders/upload",
        files={"file": ("empty.pdf", b"", "application/pdf")},
    )
    assert resp.status_code == 400


# 6. Database Core Integrity Check
def test_database_standards_integrity_after_tender_processing(db_session):
    """Ensure tender processing does not mutate or corrupt authoritative standards table."""
    standards_count = db_session.query(Standard).count()
    assert standards_count == 268  # Canonical verified count


# 7. Health Check Verification
def test_health_check_includes_module11(client):
    """Verify /api/health includes Module 11."""
    resp = client.get("/api/health")
    assert resp.status_code == 200
    assert "Module 11" in resp.json()["module"]
    assert "Tender Document Engine" in resp.json()["module"]
