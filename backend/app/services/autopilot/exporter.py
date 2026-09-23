"""
DOCX export for Autopilot tender packages: tender body, citation appendix,
red-team findings and approval block.
"""
import io
from typing import Any, Dict

from docx import Document
from docx.shared import Pt, RGBColor


def build_tender_docx(result: Dict[str, Any]) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    doc.add_heading("Tender Document — Technical Specification", level=0)
    meta = doc.add_paragraph()
    meta.add_run(f"Autopilot run {result.get('run_id', '')}  ·  Readiness: {result.get('readiness', '')}").italic = True
    doc.add_paragraph(f"Requirement: {result.get('need', '')}")

    for sec in result.get("sections", []):
        doc.add_heading(sec.get("heading", ""), level=1)
        num = sec.get("heading", "").split(".")[0]
        for i, clause in enumerate(sec.get("clauses", []), 1):
            p = doc.add_paragraph()
            p.add_run(f"{num}.{i}  ").bold = True
            p.add_run(clause.get("text", ""))
            cites = clause.get("citations") or []
            if cites:
                r = p.add_run("  [" + ", ".join(cites) + "]")
                r.font.size = Pt(8)
                r.font.color.rgb = RGBColor(0x1F, 0x4E, 0xA8)

    redteam = result.get("redteam") or {}
    if redteam.get("findings"):
        doc.add_page_break()
        doc.add_heading("Annexure A — Pre-publication Red-Team Review", level=1)
        doc.add_paragraph(
            f"Audit coverage: {redteam.get('coverage_before', 0):.0f}% before corrections, "
            f"{redteam.get('coverage_after', 0):.0f}% after. Auto-corrections applied: {redteam.get('auto_fixes', 0)}."
        )
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        for cell, head in zip(table.rows[0].cells, ["Severity", "Issue", "Resolution", "Status"]):
            cell.text = head
        for f in redteam["findings"]:
            row = table.add_row().cells
            row[0].text = f.get("severity", "")
            row[1].text = f.get("issue", "")
            row[2].text = f.get("fix", "")
            row[3].text = "Auto-fixed" if f.get("auto_fixed") else "Officer action"

    if result.get("citations"):
        doc.add_heading("Annexure B — Evidence & Citations", level=1)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        for cell, head in zip(table.rows[0].cells, ["Key", "Type", "Record", "Source dataset"]):
            cell.text = head
        for c in result["citations"]:
            row = table.add_row().cells
            row[0].text = c.get("key", "")
            row[1].text = c.get("source_type", "")
            row[2].text = c.get("details", "")
            row[3].text = c.get("source_dataset", "")

    doc.add_heading("Approval", level=1)
    doc.add_paragraph("Prepared by: ARISTEA Autopilot (decision-support draft)")
    doc.add_paragraph("Reviewed & approved by: ______________________    Designation: ____________    Date: ________")
    disclaimer = doc.add_paragraph(result.get("disclaimer", ""))
    disclaimer.runs[0].font.size = Pt(8) if disclaimer.runs else None

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
